#!/usr/bin/env python3
"""
FTA apportionment tables -> one queryable table. The data behind Ask Funding.

Each year FTA publishes ~27 Excel tables of formula apportionments. This reads the three that
carry the money transit agencies actually plan around:

  Table 3   Section 5307 + 5340   Urbanized Area Formula      (and its component breakout)
  Table 11  Section 5337          State of Good Repair
  Table 12  Section 5339          Buses and Bus Facilities

WHY THE FILES ARE PASSED IN BY HAND. transit.dot.gov returns HTTP 403 to every scripted client -
the page handler and the static /files/ path alike. We do not work around bot protection, so the
workbooks are downloaded by a person and dropped in a directory, the same arrangement the CIG
project profiles already use. FTA publishes once a year, so this is a once-a-year chore.

  python funding.py --dir /data/funding --year 2026
  python funding.py --dir /data/funding --year 2026 --db /data/funding.duckdb

Run it with --dry-run to parse and report without writing anything.
"""

import argparse
import glob
import os
import re
import sys

# ---- the shape we normalise everything into -----------------------------------------------------
# One row per (program, place, measure). Long rather than wide, because the three programs carry
# different components and a wide table would be mostly nulls - and because "how much did X get
# from Y" is the question people actually ask.
COLUMNS = [
    "fiscal_year",     # 2026
    "program",         # '5307' | '5337' | '5339'
    "uza_name",        # 'Boston, MA-NH'   (null on statewide rows)
    "state",           # 'Massachusetts'   (null where the table does not say)
    "measure",         # 'Section 5307' | 'High Intensity Fixed Guideway' | 'Apportionment' | ...
    "amount",          # whole US dollars
    "row_kind",        # 'uza' | 'state' | 'uza_total'
    "source_table",    # 'Table 3' | 'Table 11' | 'Table 12'
]

# FTA does not name the same place the same way in every table, and there is no UZA code to fall
# back on. One discrepancy exists in FY2026 and it is the largest single line in Table 11:
# $1.206bn, 27.9% of the whole High Intensity Fixed Guideway program. Left unaliased it silently
# fails to join and a quarter of 5337 disappears while every total still looks plausible.
UZA_ALIASES = {
    "New York, NY": "New York-Jersey City-Newark, NY-NJ",
}

# Rows that are arithmetic, not recipients.
TOTAL_RE = re.compile(r"^\s*(national\s+)?total\s*:?\s*$", re.I)
SUBTOTAL_RE = re.compile(r"population|statewide allocation|national total", re.I)


def canon_uza(name):
    if not name:
        return None
    n = " ".join(str(name).split())
    return UZA_ALIASES.get(n, n)


def cells(ws, row, n):
    return [ws.cell(row=row, column=c).value for c in range(1, n + 1)]


def find_header(ws, must_contain, limit=20):
    """The header row moves between files (7, 8, 9 across these three), so find it rather than
    assume it. Returns the row number, or raises - a wrong guess here corrupts everything after."""
    want = [w.lower() for w in must_contain]
    for r in range(1, min(limit, ws.max_row) + 1):
        joined = " | ".join(str(v).lower() for v in cells(ws, r, 8) if v is not None)
        if all(w in joined for w in want):
            return r
    raise ValueError("could not find a header row containing %s in sheet %r" % (must_contain, ws.title))


def money(v):
    return float(v) if isinstance(v, (int, float)) else None


# ---- Table 3: Section 5307 / 5340 ---------------------------------------------------------------
def parse_5307(path, year, report):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)

    # Two sheets carry the same money. The "Breakout" one is the better source: it splits the
    # apportionment into its four statutory components, and its UZA name sits in its own column.
    # (Note the two sheets put state and UZA in OPPOSITE columns under the same header text.)
    sheet = next((s for s in wb.sheetnames if "breakout" in s.lower()), None)
    if not sheet:
        raise ValueError("Table 3 has no Breakout sheet; got %s" % wb.sheetnames)
    ws = wb[sheet]
    hdr = find_header(ws, ["urbanized area", "uza name"])
    labels = [str(v).strip() if v else "" for v in cells(ws, hdr, 7)]
    measures = labels[2:7]          # Section 5307 | 5307 STIC | 5340 Growing States | 5340 High Density | Total

    out, skipped = [], 0
    for r in range(hdr + 1, ws.max_row + 1):
        state, uza = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
        state = " ".join(str(state).split()) if state else None
        uza = canon_uza(uza)
        # Section headers, tier subtotals and the national total all arrive with no UZA name.
        if not uza:
            skipped += 1
            continue
        # In the third section ("Amounts Apportioned to States for UZAs under 200,000") FTA puts a
        # STATE ROLLUP in the UZA column - blank col A, "Alabama" in col B - directly above that
        # state's individual small UZAs. Every real UZA name carries a ", XX" suffix and a rollup
        # does not, which is how they are told apart. Counting both double-counts the entire
        # small-UZA tier: $942,157,286 in FY2026, with every individual figure still correct.
        if "," not in uza:
            skipped += 1
            continue
        for i, m in enumerate(measures):
            if not m or m.lower() == "total":
                continue                      # Total is derivable; storing it invites double counting
            amt = money(ws.cell(row=r, column=3 + i).value)
            if amt is None:
                continue
            out.append([year, "5307", uza, state, m, amt, "uza", "Table 3"])
    report.append("  Table 3  (5307/5340): %4d rows from %s, %d non-data rows skipped"
                  % (len(out), sheet, skipped))
    return out


# ---- Table 11: Section 5337 ---------------------------------------------------------------------
def parse_5337(path, year, report):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).active
    hdr = find_header(ws, ["state", "urbanized area"])
    labels = [str(v).strip() if v else "" for v in cells(ws, hdr, 4)]
    measures = labels[2:4]          # High Intensity Fixed Guideway ... | High Intensity Motorbus ...

    out, skipped, aliased = [], 0, 0
    for r in range(hdr + 1, ws.max_row + 1):
        state, uza = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
        if not uza or (isinstance(state, str) and TOTAL_RE.match(state)):
            skipped += 1
            continue
        raw = " ".join(str(uza).split())
        uza = canon_uza(uza)
        if uza != raw:
            aliased += 1
        state = " ".join(str(state).split()) if state else None
        for i, m in enumerate(measures):
            amt = money(ws.cell(row=r, column=3 + i).value)
            if amt is None:
                continue
            out.append([year, "5337", uza, state, m, amt, "uza", "Table 11"])
    report.append("  Table 11 (5337):      %4d rows, %d non-data rows skipped, %d UZA name(s) aliased"
                  % (len(out), skipped, aliased))
    return out


# ---- Table 12: Section 5339 ---------------------------------------------------------------------
# One identifier column that means three different things depending on which section you are in.
# The merged section header is the ONLY thing that distinguishes a UZA from a state, and "New York"
# appears twice as a state. Parse this without tracking the section and states become UZAs.
SECTION_UZA = re.compile(r"apportioned to urbanized areas", re.I)
SECTION_STATE = re.compile(r"state governors|state/territory allocation", re.I)


def parse_5339(path, year, report):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).active
    hdr = find_header(ws, ["urbanized area", "apportionment"])

    out, section, skipped = [], None, 0
    counts = {"uza": 0, "state": 0}
    for r in range(hdr + 1, ws.max_row + 1):
        label, amt = ws.cell(row=r, column=1).value, money(ws.cell(row=r, column=2).value)
        if isinstance(label, str) and label.strip():
            text = " ".join(label.split())
            if SECTION_UZA.search(text):
                section = "uza"
                continue
            if SECTION_STATE.search(text):
                section = "state"
                continue
            if TOTAL_RE.match(text) or SUBTOTAL_RE.search(text):
                skipped += 1
                continue
            if amt is None or section is None:
                skipped += 1
                continue
            if section == "uza":
                out.append([year, "5339", canon_uza(text), None, "Apportionment", amt, "uza", "Table 12"])
            else:
                out.append([year, "5339", None, text, "Apportionment", amt, "state", "Table 12"])
            counts[section] += 1
        else:
            skipped += 1
    report.append("  Table 12 (5339):      %4d rows (%d UZA, %d statewide), %d non-data rows skipped"
                  % (len(out), counts["uza"], counts["state"], skipped))
    return out


PARSERS = [
    ("5307", ("table-3", "5307"), parse_5307),
    ("5337", ("table-11", "5337"), parse_5337),
    ("5339", ("table-12", "5339"), parse_5339),
]


def find_file(directory, hints):
    """Match on the FTA filename, which is not predictable enough to construct but does always
    carry the table number and the program number somewhere in it."""
    for p in sorted(glob.glob(os.path.join(directory, "*.xlsx"))):
        low = os.path.basename(p).lower()
        if any(h in low for h in hints):
            return p
    return None


# Where each workbook states its grand total, and which columns of that row we should be able to
# reproduce. Explicit per table rather than inferred: these sheets also carry population-tier
# subtotals whose labels look identical to the grand total, and summing those too reports exactly
# double the real figure - which then flags a correct parse as broken.
TOTAL_SPEC = {
    # program: (sheet picker, row label to match, columns to sum)
    "5307": ("breakout", "national total", (3, 4, 5, 6)),   # four components, NOT the Total column
    "5337": (None, "total", (3, 4)),                        # fixed guideway + motorbus
    "5339": (None, "national total", (2,)),
}


def published_total(path, program):
    """The grand total FTA printed in the workbook, so a parse can be checked against it.

    Worth the trouble: the first 5307 parser over-counted by exactly one population tier while
    every individual figure stayed correct. Only the total gave it away.
    """
    from openpyxl import load_workbook
    sheet_hint, label_want, cols = TOTAL_SPEC[program]
    wb = load_workbook(path, data_only=True)
    if sheet_hint:
        name = next((x for x in wb.sheetnames if sheet_hint in x.lower()), wb.sheetnames[0])
        ws = wb[name]
    else:
        ws = wb.active

    match = None
    for r in range(1, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if not isinstance(label, str):
            continue
        text = " ".join(label.split()).rstrip(":").lower()
        if text == label_want:
            match = r          # keep the LAST one: the grand total sits below any subtotals
    if match is None:
        return None
    vals = [money(ws.cell(row=match, column=c).value) for c in cols]
    vals = [v for v in vals if v is not None]
    return sum(vals) if vals else None


def load(directory, year):
    rows, report, missing, checks = [], [], [], []
    for program, hints, fn in PARSERS:
        path = find_file(directory, hints)
        if not path:
            missing.append("%s (looked for a filename containing %s)" % (program, " or ".join(hints)))
            continue
        got = fn(path, year, report)
        rows.extend(got)
        try:
            official = published_total(path, program)
        except Exception:
            official = None
        checks.append((program, sum(r[5] for r in got), official))
    return rows, report, missing, checks


def write_duckdb(rows, db, year):
    import duckdb
    con = duckdb.connect(db)
    con.execute("""CREATE TABLE IF NOT EXISTS fta_apportionments (
        fiscal_year INTEGER, program TEXT, uza_name TEXT, state TEXT,
        measure TEXT, amount DOUBLE, row_kind TEXT, source_table TEXT)""")
    # A re-run for the same year replaces that year rather than appending it twice.
    con.execute("DELETE FROM fta_apportionments WHERE fiscal_year = ?", [year])
    con.executemany("INSERT INTO fta_apportionments VALUES (?,?,?,?,?,?,?,?)", rows)
    con.execute("CREATE INDEX IF NOT EXISTS fta_app_uza ON fta_apportionments (uza_name)")
    con.execute("CREATE INDEX IF NOT EXISTS fta_app_prog ON fta_apportionments (program, fiscal_year)")
    n = con.execute("SELECT count(*) FROM fta_apportionments WHERE fiscal_year = ?", [year]).fetchone()[0]
    con.close()
    return n


def main():
    ap = argparse.ArgumentParser(description="Load FTA apportionment tables into DuckDB.")
    ap.add_argument("--dir", default=os.environ.get("FUNDING_DIR", "/data/funding"),
                    help="directory holding the downloaded .xlsx tables")
    ap.add_argument("--db", default=os.environ.get("FUNDING_DB", "/data/funding.duckdb"))
    ap.add_argument("--year", type=int, required=True, help="the FTA fiscal year these tables are for")
    ap.add_argument("--dry-run", action="store_true", help="parse and report, write nothing")
    a = ap.parse_args()

    if not os.path.isdir(a.dir):
        raise SystemExit("No such directory: %s\nDownload the tables by hand and drop them there "
                         "(transit.dot.gov blocks automated fetching)." % a.dir)

    rows, report, missing, checks = load(a.dir, a.year)
    print("FY %s apportionments from %s" % (a.year, a.dir))
    for line in report:
        print(line)
    for m in missing:
        print("  MISSING: %s" % m, file=sys.stderr)
    if not rows:
        raise SystemExit("Nothing parsed; refusing to write.")

    print("")
    print("  reconciliation against the totals FTA printed in each workbook:")
    bad = []
    for program, parsed, official in checks:
        if official is None:
            print("    Section %-5s parsed $%-18s (no published total to check against)"
                  % (program, format(int(parsed), ",")))
            continue
        delta = parsed - official
        flag = "OK" if abs(delta) < 1 else ("MISMATCH %+d" % round(delta))
        print("    Section %-5s parsed $%-18s published $%-18s %s"
              % (program, format(int(parsed), ","), format(int(official), ","), flag))
        if abs(delta) >= 1:
            bad.append(program)
    if bad:
        raise SystemExit("Refusing to write: " + ", ".join(bad) + " does not reconcile to the "
                         "published total. A parse that is off by a whole tier still looks "
                         "plausible row by row, so this is checked rather than trusted.")

    if a.dry_run:
        print("\n  --dry-run: nothing written")
        return
    n = write_duckdb(rows, a.db, a.year)
    print("\n  wrote %d rows to %s" % (n, a.db))


if __name__ == "__main__":
    main()
